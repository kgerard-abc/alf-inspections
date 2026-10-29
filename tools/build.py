"""Check data/facilities.json and write it into index.html.

  python3 tools/build.py            check, then update index.html (and its "Data last updated" date if the data changed)
  python3 tools/build.py --check    check only; fails if index.html doesn't match the data
  python3 tools/build.py --xlsx PATH   also write the master spreadsheet

Errors stop the build. Warnings are printed but don't stop it.
"""
import argparse
import datetime
import json
import re
import sys
from zoneinfo import ZoneInfo

from alf import (DATA, INDEX, NO_FINDINGS, REPORT_KEYS, SCRIPT_RE, TIER_RANK, TYPES,
                 compact, file_id, load, recompute)

UPDATED_RE = re.compile(r"(Data last updated <strong>)([^<]+)(</strong>)")


def validate(facilities):
    errors, warns = [], []
    seen_ids = {}
    keys = set()
    blank = 0
    for f in facilities:
        name = f.get("name", "?")
        if not f.get("key") or f["key"] in keys:
            errors.append(f"{name}: missing or repeated facility key")
        keys.add(f.get("key"))
        for r in f.get("reports", []):
            where = f"{name} {r.get('date')}"
            if set(r) != REPORT_KEYS:
                errors.append(f"{where}: fields are {sorted(r)}, expected {sorted(REPORT_KEYS)}")
                continue
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", r["date"] or ""):
                errors.append(f"{where}: bad date")
            fid = file_id(r["link"])
            if not fid:
                errors.append(f"{where}: link has no Drive file ID")
            elif fid in seen_ids:
                errors.append(f"{where}: same report file as {seen_ids[fid]} (duplicate)")
            else:
                seen_ids[fid] = where
            if r["no_findings"]:
                if r["tier"] is not None or r["paragraphs"] or r["headline"] != NO_FINDINGS:
                    errors.append(f"{where}: marked no findings but has a severity or findings")
            else:
                if not r["paragraphs"]:
                    errors.append(f"{where}: has findings but no finding text")
                    continue
                bad = [p for p in r["paragraphs"] if p.get("tier") not in TIER_RANK]
                if bad:
                    errors.append(f"{where}: a finding has an unknown severity")
                    continue
                top = max(r["paragraphs"], key=lambda p: TIER_RANK[p["tier"]])["tier"]
                if r["tier"] != top:
                    errors.append(f"{where}: report severity {r['tier']} but most severe finding is {top}")
                if not r["headline"].strip():
                    blank += 1
            if r["type"] not in TYPES + ["Inspection"]:
                warns.append(f"{where}: unusual type shown on the site: {r['type']!r}")
    if blank:
        warns.append(f"{blank} older reports have findings but a blank summary")
    return errors, warns


def write_xlsx(facilities, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    head = ["Facility", "License no.", "City", "Exit date", "Type", "Severity", "No findings",
            "Summary", "Findings", "Report link", "RIDOH file ID"]
    ws.append(head)
    for f in sorted(facilities, key=lambda x: x["name"]):
        for r in f["reports"]:
            ws.append([f["name"], f.get("license_no"), f.get("city"), r["date"], r["type"],
                       (r["tier"] or "").replace("_", " "), "yes" if r["no_findings"] else "",
                       r["headline"], "\n\n".join(p["text"] for p in r["paragraphs"]),
                       r["link"], file_id(r["link"])])
    ws2 = wb.create_sheet("Facilities")
    cols = ["name", "license_no", "address", "city", "zip", "phone", "capacity", "badge",
            "report_count", "most_recent_date", "website"]
    ws2.append(cols)
    for f in sorted(facilities, key=lambda x: x["name"]):
        ws2.append([f.get(c) for c in cols])
    for sheet in (ws, ws2):
        for c in sheet[1]:
            c.font = Font(name="Helvetica", bold=True)
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Helvetica")
                c.alignment = Alignment(vertical="top", wrap_text=False)
    for col, w in zip("ABCDEFGHIJK", [34, 12, 16, 11, 30, 13, 11, 50, 60, 40, 36]):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--xlsx")
    args = ap.parse_args()

    facilities = recompute(load(DATA))
    errors, warns = validate(facilities)
    for w in warns:
        print("warning:", w)
    for e in errors:
        print("ERROR:", e)
    print(f"{len(facilities)} facilities, {sum(len(f['reports']) for f in facilities)} reports, "
          f"{len(errors)} error(s), {len(warns)} warning(s)")
    if errors:
        sys.exit(1)

    html = INDEX.read_text(encoding="utf-8")
    m = SCRIPT_RE.search(html)
    new_json = compact(facilities)
    changed = m.group(2) != new_json
    # Badges are a snapshot: they drift as reports pass the 3-year mark, and the page recomputes
    # them live anyway. Only a change in the actual data counts as "changed" (and bumps the date).
    def no_badges(text):
        return [{k: v for k, v in f.items() if k != "badge"} for f in json.loads(text)]
    content_changed = changed and no_badges(m.group(2)) != no_badges(new_json)
    if args.check:
        if content_changed:
            print("ERROR: index.html does not match data/facilities.json (run python3 tools/build.py)")
            sys.exit(1)
        print("index.html matches the data")
    elif changed and not content_changed:
        INDEX.write_text(html[:m.start(2)] + new_json + html[m.end(2):], encoding="utf-8")
        print("index.html: badges refreshed for the 3-year window; 'Data last updated' left alone")
    elif changed:
        html = html[:m.start(2)] + new_json + html[m.end(2):]
        today = datetime.datetime.now(ZoneInfo("America/New_York")).date()
        stamp = f"{today:%B} {today.day}, {today.year}"
        html, n = UPDATED_RE.subn(lambda mm: mm.group(1) + stamp + mm.group(3), html)
        if n != 1:
            print("warning: 'Data last updated' line not found")
        INDEX.write_text(html, encoding="utf-8")
        print(f"index.html updated; 'Data last updated' set to {stamp}")
    else:
        print("index.html already matches the data")
    if args.xlsx:
        write_xlsx(facilities, args.xlsx)
        print("spreadsheet written:", args.xlsx)


if __name__ == "__main__":
    main()
